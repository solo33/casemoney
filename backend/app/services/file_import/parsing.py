"""import_csv: parsing."""
from __future__ import annotations

import csv
import io
from html.parser import HTMLParser
from datetime import datetime, date
from pathlib import Path
from typing import Optional
from app.services.file_import.types import COLUMN_ALIASES, POSITIONAL_COLUMNS, ParsedRow, REQUIRED_COLUMNS, SPLIT_AMOUNT_COLUMNS


def _stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _parse_amount(value) -> float:
    """Преобразует '-600,00' → -600.0"""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _parse_date(value) -> Optional[str]:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    s = value.strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _normalize_header(raw_header: list) -> dict[str, int]:
    normalized = {}
    for idx, value in enumerate(raw_header):
        key = _stringify(value).lower().strip()
        key = COLUMN_ALIASES.get(key, key)
        if key and key not in normalized:
            normalized[key] = idx
    return normalized


def _parse_rows(rows: list[list]) -> list[ParsedRow]:
    if not rows:
        return []

    header_map = _normalize_header(rows[0])
    header_keys = set(header_map)
    has_header = REQUIRED_COLUMNS.issubset(header_keys) or SPLIT_AMOUNT_COLUMNS.issubset(header_keys)
    if not has_header:
        header_map = {name: idx for idx, name in enumerate(POSITIONAL_COLUMNS)}

    parsed: list[ParsedRow] = []
    start = 1 if has_header else 0
    for raw_idx, raw in enumerate(rows[start:], start=start + 1):
        if not raw or all(not _stringify(c) for c in raw):
            continue

        def get(col: str):
            idx = header_map.get(col)
            if idx is None or idx >= len(raw):
                return ""
            return raw[idx]

        date_iso = _parse_date(get("date"))
        account = _stringify(get("account"))
        raw_record = _stringify(get("record"))
        category_path = _stringify(get("category"))
        amount = _parse_amount(get("amount"))
        if not amount and ("expense" in header_map or "income" in header_map):
            expense = _parse_amount(get("expense"))
            income = _parse_amount(get("income"))
            amount = income if income else expense
        currency = (_stringify(get("currency")) or "RUB").upper()
        description = _stringify(get("description")) or None
        transfer_to = _stringify(get("transfer")) or None

        record_lower = raw_record.lower()
        if not category_path and raw_record:
            if record_lower.startswith("расход:"):
                category_path = raw_record.split(":", 1)[1].strip()
            elif record_lower.startswith("доход:"):
                category_path = raw_record.split(":", 1)[1].strip()
        if not transfer_to and record_lower.startswith("перевод"):
            if ":" in raw_record:
                transfer_to = raw_record.split(":", 1)[1].strip()
            else:
                transfer_to = raw_record.replace("Перевод", "").replace("перевод", "").strip()

        if transfer_to:
            tx_type = "transfer"
        elif record_lower.startswith("доход:"):
            tx_type = "income"
            amount = abs(amount)
        elif record_lower.startswith("расход:"):
            tx_type = "expense"
            amount = -abs(amount)
        else:
            tx_type = "expense" if amount < 0 else ("income" if amount > 0 else "expense")

        parent = child = None
        if category_path:
            parts = [p.strip() for p in category_path.split("\\") if p.strip()]
            if len(parts) >= 2:
                parent, child = parts[0], parts[-1]
            elif len(parts) == 1:
                parent = parts[0]

        err = None
        if not date_iso:
            err = "не удалось распарсить дату"
        elif not account:
            err = "пустой счет"

        parsed.append(ParsedRow(
            line_no=raw_idx,
            date=date_iso,
            account=account,
            category_path=category_path or None,
            category_parent=parent,
            category_child=child,
            amount=amount,
            abs_amount=abs(amount),
            currency=currency,
            description=description,
            transfer_to=transfer_to,
            tx_type=tx_type,
            error=err,
        ))
    return _merge_mirrored_transfers(parsed)


def _merge_mirrored_transfers(rows: list[ParsedRow]) -> list[ParsedRow]:
    """Некоторые экспорты (напр. HomeMoney) кладут один перевод двумя строками —
    расход на счёте-источнике и доход на счёте-получателе, — причём при обмене
    валюты суммы и валюты в этих строках отличаются (курс на момент перевода).
    Обе строки описывают одно и то же движение денег; без объединения перевод
    применяется дважды, а без учёта расхождения валют зачисление считается в
    валюте списания. Здесь исходящая строка (amount < 0) ищет зеркальную
    входящую (amount > 0, тот же день, встречные account/transfer_to) и
    забирает её валюту/сумму в to_currency/to_amount; входящая строка-зеркало
    выбрасывается. Строки без пары (одна нога перевода в исходных данных)
    остаются как есть — зачисление считается в валюте списания, как раньше."""
    def match(rows_idx: list[int], key_fn, drop: set[int], matched: set[int]):
        """Один проход сопоставления. key_fn(row) → (my_key, mirror_key)."""
        pending_out: dict[tuple, list[int]] = {}
        pending_in: dict[tuple, list[int]] = {}
        for i in rows_idx:
            r = rows[i]
            my_key, mirror_key = key_fn(r)
            if r.amount < 0:
                waiting = pending_in.get(mirror_key)
                if waiting:
                    in_idx = waiting.pop(0)
                    if not waiting:
                        del pending_in[mirror_key]
                    in_row = rows[in_idx]
                    r.to_currency = in_row.currency
                    r.to_amount = in_row.abs_amount
                    drop.add(in_idx)
                    matched.add(i)
                    matched.add(in_idx)
                else:
                    pending_out.setdefault(my_key, []).append(i)
            else:
                waiting = pending_out.get(mirror_key)
                if waiting:
                    out_idx = waiting.pop(0)
                    if not waiting:
                        del pending_out[mirror_key]
                    out_row = rows[out_idx]
                    out_row.to_currency = r.currency
                    out_row.to_amount = r.abs_amount
                    drop.add(i)
                    matched.add(i)
                    matched.add(out_idx)
                else:
                    pending_in.setdefault(my_key, []).append(i)

    transfer_idx = [
        i for i, r in enumerate(rows)
        if not r.error and r.transfer_to and r.tx_type == "transfer"
    ]
    drop: set[int] = set()
    matched: set[int] = set()

    # Проход 1 — строгий: та же валюта и сумма. Не даёт склеить два РАЗНЫХ
    # встречных перевода одного дня (A→B 100 и B→A 50 останутся раздельными).
    def strict_key(r):
        acc = r.account.strip().lower()
        to = r.transfer_to.strip().lower()
        amt = round(r.abs_amount, 2)
        return (r.date, acc, to, r.currency, amt), (r.date, to, acc, r.currency, amt)

    match(transfer_idx, strict_key, drop, matched)

    # Проход 2 — по остатку: без валюты/суммы (перевод с конвертацией — суммы
    # в ногах разные). Экспорты HomeMoney записывают зеркальные ноги в одном
    # порядке, поэтому несколько обменов между теми же счетами за день можно
    # детерминированно сопоставить по очереди. Лишняя строка остаётся отдельным
    # односторонним переводом — это уже поддерживаемое поведение импортёра.
    def loose_key(r):
        acc = r.account.strip().lower()
        to = r.transfer_to.strip().lower()
        return (r.date, acc, to)

    pending_out: dict[tuple, list[int]] = {}
    pending_in: dict[tuple, list[int]] = {}
    for i in transfer_idx:
        if i in matched:
            continue
        r = rows[i]
        bucket = pending_out if r.amount < 0 else pending_in
        bucket.setdefault(loose_key(r), []).append(i)

    for key, out_indices in pending_out.items():
        date, acc, to = key
        mirror = (date, to, acc)
        in_indices = pending_in.get(mirror, [])
        for out_idx, in_idx in zip(out_indices, in_indices):
            rows[out_idx].to_currency = rows[in_idx].currency
            rows[out_idx].to_amount = rows[in_idx].abs_amount
            drop.add(in_idx)
            matched.update((out_idx, in_idx))

    return [r for i, r in enumerate(rows) if i not in drop]


def _repair_wrapped_csv_descriptions(rows: list[list]) -> None:
    """Восстановить перенос строки в некавычённом поле description.

    Некоторые старые экспорты HomeMoney содержат физический перевод строки
    внутри описания без обязательных CSV-кавычек. ``csv.reader`` видит хвост
    такого описания отдельной строкой вида ``["Яйца", ""]``. Оставляем строку
    на месте (чтобы номера последующих строк не сдвигались), переносим текст в
    описание предыдущей операции и очищаем ложную строку.
    """
    if len(rows) < 3:
        return

    header_map = _normalize_header(rows[0])
    if not REQUIRED_COLUMNS.issubset(set(header_map)):
        return
    date_idx = header_map["date"]
    description_idx = header_map.get("description")
    if description_idx is None:
        return

    for idx in range(1, len(rows)):
        raw = rows[idx]
        non_empty = [(col, _stringify(value)) for col, value in enumerate(raw) if _stringify(value)]
        if len(non_empty) != 1 or non_empty[0][0] != date_idx or _parse_date(raw[date_idx]):
            continue

        previous = rows[idx - 1]
        if date_idx >= len(previous) or not _parse_date(previous[date_idx]):
            continue

        while len(previous) <= description_idx:
            previous.append("")
        continuation = non_empty[0][1]
        previous_description = _stringify(previous[description_idx])
        previous[description_idx] = " ".join(part for part in (previous_description, continuation) if part)
        rows[idx] = [""] * len(raw)


def parse_csv(content: bytes) -> list[ParsedRow]:
    """Parse raw CSV content into ParsedRow objects."""
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Не удалось распознать кодировку файла")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";"
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    _repair_wrapped_csv_descriptions(rows)
    return _parse_rows(rows)


def parse_xlsx(content: bytes) -> list[ParsedRow]:
    """Parse XLSX content using the first worksheet."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ValueError("Для импорта XLSX установите зависимость openpyxl") from e

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    return _parse_rows(rows)


def parse_xls(content: bytes) -> list[ParsedRow]:
    """Parse legacy XLS content using the first worksheet."""
    prefix = content[:256].lstrip().lower()
    if prefix.startswith(b"<") or prefix.startswith(b"\xef\xbb\xbf<"):
        return parse_html_table(content)

    try:
        import xlrd
    except ImportError as e:
        raise ValueError("Для импорта XLS установите зависимость xlrd") from e

    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    rows = []
    for row_idx in range(sheet.nrows):
        row = []
        for col_idx in range(sheet.ncols):
            cell = sheet.cell(row_idx, col_idx)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = datetime(*xlrd.xldate_as_tuple(value, book.datemode))
            row.append(value)
        rows.append(row)
    return _parse_rows(rows)


def parse_html_table(content: bytes) -> list[ParsedRow]:
    """Parse bank exports saved as .xls but containing an HTML table."""
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Не удалось распознать кодировку HTML-таблицы")

    parser = _SimpleTableParser()
    parser.feed(text)
    rows = parser.rows
    if not rows:
        return []
    return _parse_rows(rows)


class _SimpleTableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._in_cell = False

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
            self._in_cell = True

    def handle_data(self, data):
        if self._in_cell and self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            value = " ".join("".join(self._cell).split())
            self._row.append(value)
            self._cell = None
            self._in_cell = False
        elif tag == "tr" and self._row is not None:
            if any(str(c).strip() for c in self._row):
                self.rows.append(self._row)
            self._row = None


def parse_file(filename: str, content: bytes) -> list[ParsedRow]:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv" or not suffix:
        return parse_csv(content)
    if suffix == ".xlsx":
        return parse_xlsx(content)
    if suffix == ".xls":
        return parse_xls(content)
    raise ValueError("Поддерживаются файлы CSV, XLSX и XLS")
